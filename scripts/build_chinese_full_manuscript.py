#!/usr/bin/env python3
"""Build a Chinese full-text DOCX from the reviewed Chinese Markdown draft."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from docx_style import add_heading as add_styled_heading
from docx_style import add_paragraph, set_run_font


def add_heading(document: Document, text: str, level: int) -> None:
    add_styled_heading(
        document,
        text,
        size=14 if level == 1 else 12,
        space_before=8 if level == 1 else 6,
    )


def add_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(value)
            set_run_font(run, 8.5, bold=row_index == 0)
    document.add_paragraph()


def add_page_number(document: Document) -> None:
    for section in document.sections:
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = " PAGE "
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.extend([begin, instruction, end])


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw = lines[index].strip().strip("|")
        cells = [cell.strip() for cell in raw.split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--markdown", required=True)
    parser.add_argument("--figure", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    lines = Path(args.markdown).read_text(encoding="utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    add_page_number(document)

    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            if rows:
                add_table(document, rows)
            continue
        if stripped.startswith("### "):
            add_heading(document, stripped[4:], 2)
        elif stripped.startswith("## "):
            add_heading(document, stripped[3:], 1)
        elif stripped.startswith("# "):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(stripped[2:])
            set_run_font(run, 18, True)
        elif stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            run = paragraph.add_run(stripped[2:])
            set_run_font(run, 10.5)
        else:
            add_paragraph(document, stripped)
        index += 1

    document.add_paragraph()
    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(Path(args.figure)), width=Inches(6.8))

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)


if __name__ == "__main__":
    main()
