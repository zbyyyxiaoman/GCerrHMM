#!/usr/bin/env python3
"""Create the W3.8 review copy with Figures 1-9 embedded after captions."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph


FIGURES = {
    "Figure 1.": "Figure1_gc_context.png",
    "Figure 2.": "Figure2_simulator_landscape.png",
    "Figure 3.": "Figure3_context_layers.png",
    "Figure 4.": "Figure4_overview_innovation.png",
    "Figure 5.": "Figure5_reads_cross_tool.png",
    "Figure 6.": "Figure6_alignment.png",
    "Figure 7.": "Figure7_variant_calling.png",
    "Figure 8.": "Figure8_assembly_phasing.png",
    "Figure 9.": "Figure9_ablation.png",
}


def insert_picture_after(
    anchor: Paragraph,
    image: Path,
    width_inches: float = 6.2,
) -> Paragraph:
    new_element = OxmlElement("w:p")
    anchor._p.addnext(new_element)
    paragraph = Paragraph(new_element, anchor._parent)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Inches(0)
    paragraph.paragraph_format.space_after = Inches(0.12)
    run = paragraph.add_run()
    run.add_picture(str(image), width=Inches(width_inches))
    return paragraph


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--figure-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    document = Document(args.input)
    captions = {
        prefix: paragraph
        for paragraph in document.paragraphs
        for prefix in FIGURES
        if paragraph.text.strip().startswith(prefix)
    }
    missing = sorted(set(FIGURES) - set(captions))
    if missing:
        raise RuntimeError(f"missing figure captions: {missing}")

    for prefix, filename in FIGURES.items():
        image = args.figure_dir / filename
        if not image.exists():
            raise FileNotFoundError(image)
        caption = captions[prefix]
        caption.paragraph_format.keep_with_next = True
        insert_picture_after(caption, image)

    document.core_properties.comments = (
        "W3.8 review copy with Figures 1-9 embedded after their legends."
    )
    document.core_properties.last_modified_by = "GCerrHMM project"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
