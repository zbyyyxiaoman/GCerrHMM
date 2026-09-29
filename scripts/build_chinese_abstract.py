#!/usr/bin/env python3
"""Build a Chinese figure-inclusive abstract for internal review."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from docx_style import add_heading, add_paragraph, set_run_font

def add_callout(document: Document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    set_run_font(run, 11, True)
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), "EAF2F8")
    cell._tc.get_or_add_tcPr().append(fill)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--contact-sheet", required=True)
    parser.add_argument("--output-docx", required=True)
    args = parser.parse_args()

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "GCerrHMM：面向长读长测序模拟的 GC 感知错误隐马尔可夫模型"
        "与可复现评估框架"
    )
    set_run_font(run, 17, True)
    add_heading(document, "背景")
    add_paragraph(
        document,
        "长读长测序已广泛用于结构变异检测、单倍型组装和复杂基因组研究。"
        "当真实真值不可得时，模拟读长被用来检验比对、变异检测和组装流程，"
        "因此模拟器的错误过程直接决定下游基准的可信度。现有模拟器多从真实"
        "比对中估计全局错误或质量分布，但真实 ONT 错误会随局部序列上下文"
        "变化。问题不仅在于模型是否条件化，还在于评估指标是否真的能看见"
        "这种条件化。",
    )

    add_heading(document, "结果")
    add_paragraph(
        document,
        "我们提出 GCerrHMM：一个从真实 ONT 比对训练的状态转移错误 HMM。"
        "其转移矩阵按 100 bp 窗口的局部 GC bin 条件化，替换矩阵保持全局；"
        "长度感知的 indel 状态用于减少固定长度插入偏置。评估采用预注册的 "
        "GC 分层曲线、六基因组方向性面板和覆盖度匹配的人类 chr21 30x 下游"
        "面板。",
    )
    add_paragraph(
        document,
        "在六个基因组中，GC-aware 模型在四个基因组的两个重复中方向性优于"
        "1-bin 对照；其中 H. sapiens chr21 的方向性差异最大"
        "（B r = -0.579/-0.091，C r = 0.917/0.861）。预注册的"
        "双物种 × 双判据 gate 为跨物种普遍性设定了判定范围；"
        "本研究的重点是可重复的方向性效应及其测量边界，"
        "跨物种普遍性由预注册 gate 界定。",
    )
    add_paragraph(
        document,
        "不确定性被拆分为两类：between-window bootstrap interval 宽约 1，"
        "反映窗口层面的结构性不确定性；read-sampling SD 在 5x/10x/15x"
        "为 0.176/0.194/0.009（6/3/2 个 subsets），其中 15x 由两个"
        "互补子集估计。复合分布分数用于总结整体分布保真度，"
        "GC 条件化由专门的 GC 分层指标评估。chr21 30x 面板中，GCerrHMM 在 "
        "alignment identity 上 rank 2、assembly Borda 上 rank 1，"
        "支持相对于真实 anchor 的下游一致性。",
    )

    add_heading(document, "结论")
    add_paragraph(
        document,
        "GCerrHMM 提供了一个可训练、可复现、能显式表达局部 GC 条件化错误"
        "状态的模拟框架，并给出了方向性 GC 证据和下游一致性结果。当前"
        "证据支持在 GC 结构复杂基因组中的方向性结论；更大规模的预注册 "
        "panel 将检验跨物种普遍性。本文的核心方法学结论是，模拟器评估"
        "必须让评估指标与被检验的模型性质对齐：包含 GC 条件的指标用于"
        "检验 GC 建模，描述性复合分数用于总结整体分布保真度。",
    )

    add_heading(document, "图 1-图 9 总览")
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(
        str(Path(args.contact_sheet)),
        width=Inches(6.8),
    )
    captions = [
        "图 1：六基因组局部 GC 结构与真实/模拟 GC-error 曲线。",
        "图 2：2017-2026 长读长模拟器与基准论文景观。",
        "图 3：homopolymer 上下文与评估层级。",
        "图 4：GCerrHMM 整体流程与模型结构。",
        "图 5：reads-level 分布保真度。",
        "图 6：alignment-level 对比。",
        "图 7：variant calling 对比。",
        "图 8：assembly/phasing 与真实 anchor。",
        "图 9：GC-bin、覆盖度和状态空间消融。",
    ]
    for caption in captions:
        add_paragraph(document, caption, size=9.5)

    output = Path(args.output_docx)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)


if __name__ == "__main__":
    main()
