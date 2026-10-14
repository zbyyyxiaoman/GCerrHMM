#!/usr/bin/env python3
"""Add phase-2 A. thaliana replication data to Additional file 3."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def add_header(sheet, headers: list[str]) -> None:
    sheet.append(headers)
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")


def autosize(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--base", type=Path, required=True)
    parser.add_argument("--stats", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    payload = json.loads(args.stats.read_text())
    second = payload["second_cohort"]
    first = payload["first_cohort"]
    discovery = payload["discovery_cohort"]
    pooled = payload["all_seed_sensitivity"]

    workbook = load_workbook(args.base)
    for name in ("Athaliana_replication_phase2", "Athaliana_replication_summary"):
        if name in workbook.sheetnames:
            del workbook[name]

    sheet = workbook.create_sheet("Athaliana_replication_phase2")
    add_header(sheet, ["cohort", "species", "gc_bins", "seed", "gc_pearson_r", "delta_vs_1bin"])
    for seed in second["seeds"]:
        seed_key = str(seed)
        r5 = float(second["r_5bin"][seed_key])
        r1 = float(second["r_1bin"][seed_key])
        delta = float(second["delta"][seed_key])
        sheet.append(["phase2", "A. thaliana", 5, seed, r5, delta])
        sheet.append(["phase2", "A. thaliana", 1, seed, r1, ""])

    summary = workbook.create_sheet("Athaliana_replication_summary")
    add_header(
        summary,
        [
            "cohort",
            "seeds",
            "n",
            "mean_delta_r",
            "sd_delta_r",
            "ci95_low",
            "ci95_high",
            "paired_t_p",
            "supported",
        ],
    )
    for name, values in (
        ("discovery", discovery),
        ("first_replication", first),
        ("second_replication", second),
        ("pooled_12_seed", pooled),
    ):
        summary.append(
            [
                name,
                ",".join(str(seed) for seed in values["seeds"]),
                values["n"],
                values["mean_delta"],
                values["sd_delta"],
                values["ci95_low"],
                values["ci95_high"],
                values["paired_t_p"],
                values["supported"],
            ]
        )

    autosize(sheet, [14, 14, 10, 10, 16, 18])
    autosize(summary, [24, 24, 8, 14, 14, 14, 14, 14, 12])
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
