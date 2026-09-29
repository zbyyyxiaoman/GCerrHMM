#!/usr/bin/env python3
"""Validate the BMC submission-format manuscript and its figure package."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from build_w3_submission import FIGURE_MAP


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", required=True)
    parser.add_argument("--figures", required=True)
    args = parser.parse_args()

    document = Document(Path(args.docx))
    errors: list[str] = []
    warnings: list[str] = []

    if any("w:drawing" in paragraph._p.xml for paragraph in document.paragraphs):
        errors.append("submission manuscript still contains embedded images")
    if any(
        "Background Figure" in paragraph.text
        for paragraph in document.paragraphs
    ):
        errors.append("old Background Figure labels remain")
    if any(
        br.get(qn("w:type")) == "page"
        for br in document.element.body.iter(qn("w:br"))
    ):
        errors.append("explicit page break remains in manuscript")

    text_parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            text_parts.extend(cell.text for cell in row.cells)
    text = "\n".join(text_parts)
    if (
        "W3 submission-format manuscript" not in text
        and "W3.1 submission-format manuscript" not in text
        and "W3.2 submission-format manuscript" not in text
        and "W3.3 submission-format manuscript" not in text
    ):
        errors.append("title-page status line was not updated")
    if "Additional files" not in text:
        errors.append("Additional files section is missing")
    if "Supplementary material" in text:
        errors.append("old Supplementary material heading remains")
    for forbidden in (
        "emission model is conditioned",
        "undefined correlation",
        "exactly what the mechanism predicts",
        "Simulated-read phasing results",
        "Fig. 4c now lists",
        "A boundary condition:",
        "Ecoli |",
        "Scerevisiae |",
        "Hsapiens_chr21 |",
    ):
        if forbidden in text:
            errors.append(f"stale or inconsistent text remains: {forbidden}")
    for required in (
        "state-transition model is conditioned",
        "A one-bin model cannot generate",
        "A testable boundary hypothesis",
        "Real 30x ONT anchor",
        "R5 Borda is the mean",
        "Additional_file_1_species_panel.xlsx",
    ):
        if required not in text:
            errors.append(f"W3.1 correction missing: {required}")

    for number in range(1, 10):
        if f"Figure {number}." not in text:
            errors.append(f"missing legend for Figure {number}")
        if f"Fig. {number}" not in text:
            errors.append(f"missing in-text citation for Fig. {number}")

    reference_labels = re.findall(r"(?m)^\d+\.\s", text)
    if len(reference_labels) < 16:
        errors.append(
            f"expected at least 16 numbered references, found "
            f"{len(reference_labels)}"
        )

    for section in document.sections:
        if section._sectPr.find(qn("w:lnNumType")) is None:
            errors.append("section is missing line numbering")
        if "PAGE" not in section.footer._element.xml:
            errors.append("section footer is missing page numbering")

    shaded_cells = 0
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                shaded_cells += len(list(cell._tc.iter(qn("w:shd"))))
    if shaded_cells:
        errors.append(f"table shading remains in {shaded_cells} cells")

    figure_dir = Path(args.figures)
    expected_stems = list(FIGURE_MAP.values())
    for stem in expected_stems:
        for suffix in (".png", ".pdf"):
            if not (figure_dir / f"{stem}{suffix}").is_file():
                errors.append(f"missing figure file {stem}{suffix}")

    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name not in {"Title", "Heading 1", "Heading 2", "Heading 3"}:
            if paragraph.text.strip() and paragraph.paragraph_format.line_spacing != 2.0:
                warnings.append(
                    f"non-double spacing: {paragraph.text.strip()[:60]}"
                )

    if errors:
        print("W3_AUDIT_FAILED")
        for error in errors:
            print(f"ERROR={error}")
    else:
        print("W3_AUDIT_OK")
    for warning in warnings[:10]:
        print(f"WARNING={warning}")
    if warnings:
        print(f"WARNING_COUNT={len(warnings)}")


if __name__ == "__main__":
    main()
