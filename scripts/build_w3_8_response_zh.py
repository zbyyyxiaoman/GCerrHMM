#!/usr/bin/env python3
"""Create the plain Chinese W3.8 reply from the W3.7 review note."""

from __future__ import annotations

import argparse
import json
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


def set_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def find(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(
            f"expected one paragraph starting with {prefix!r}: {len(matches)}"
        )
    return matches[0]


def insert_before(anchor: Paragraph, text: str, template: Paragraph) -> Paragraph:
    new_element = deepcopy(template._p)
    anchor._p.addprevious(new_element)
    paragraph = Paragraph(new_element, anchor._parent)
    set_text(paragraph, text)
    return paragraph


def format_p(value: float) -> str:
    value = float(value)
    return "<0.0001" if value < 0.0001 else f"{value:.4f}"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--base", type=Path, required=True)
    parser.add_argument("--stats", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    stats = json.loads(args.stats.read_text(encoding="utf-8"))
    second = stats["second_cohort"]
    pooled = stats["all_seed_sensitivity"]
    supported = bool(second["supported"])

    document = Document(args.base)
    for paragraph in document.paragraphs:
        text = paragraph.text.replace("W3.7", "W3.8")
        set_text(paragraph, text)
        if "CycSim" in paragraph.text:
            set_text(
                paragraph,
                "CycSim 是上下文感知长读长模拟器，与本文直接相关。新稿已在 "
                "Background、Discussion 和 Table 10 中明确写出：CycSim 学习"
                "序列上下文依赖的错误谱，主要服务于 mapping 参数优化；"
                "GCerrHMM 条件化的是 GC-bin 的 error-state transition，并用"
                "GC 分层曲线和预注册规则评价。由于本文的跨工具面板已经在统一 "
                "coverage、profile 和 truth-injection 协议下冻结，我们没有把 "
                "CycSim 事后加入旧面板做非等价头对头，而是保留设计层比较，"
                "并把匹配实验列为后续工作。",
            )

    summary = find(document, "四、简单总结")
    set_text(summary, "五、简单总结")
    heading_template = find(document, "二、没有改的地方")
    body_template = find(document, "没有改的主要是预注册 gate")

    insert_before(
        summary,
        "四、第二轮独立复现结果",
        heading_template,
    )
    if supported:
        result_sentence = (
            "因此，这个在运行前锁定终点和分析规则的队列，可以支持拟南芥中的"
            "单物种确认性结果。"
        )
    else:
        result_sentence = (
            "因此，这个队列仍未达到预设的确认阈值；我们按原规则报告为未确认，"
            "不把方向性结果写成确认性结论。"
        )
    insert_before(
        summary,
        "我们又按预先写好的方案运行了第二轮拟南芥独立队列，使用 6 个全新 "
        "seed（77、88、99、111、122、133），只比较 1-bin 与 5-bin。第二轮"
        f"队列的平均 delta r = {float(second['mean_delta']):.4f}，95% CI "
        f"{float(second['ci95_low']):.4f} 到 {float(second['ci95_high']):.4f}，"
        f"配对 t 检验 p = {float(second['paired_t_p']):.4f}。"
        + result_sentence,
        body_template,
    )
    insert_before(
        summary,
        "把 3 个 discovery seed、第一轮 3 个 seed 和第二轮 6 个 seed 合并后，"
        f"12 个 seed 的平均 delta r = {float(pooled['mean_delta']):.4f}，"
        f"95% CI {float(pooled['ci95_low']):.4f} 到 "
        f"{float(pooled['ci95_high']):.4f}，p "
        f"{format_p(pooled['paired_t_p'])}。这个合并结果只作为敏感性分析；"
        "原来的六物种 gate 仍然不追溯修改。",
        body_template,
    )

    document.core_properties.comments = (
        "W3.8 中文修改回复；加入第二轮独立复现实验和结果。"
    )
    document.core_properties.last_modified_by = "GCerrHMM project"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
