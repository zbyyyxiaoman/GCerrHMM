"""Shared DOCX paragraph and run styles for the manuscript builders."""

from __future__ import annotations

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def set_run_font(run, size: float, bold: bool = False) -> None:
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.bold = bold
    properties = run._element.get_or_add_rPr()
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Microsoft YaHei")
    fonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_paragraph(
    document,
    text: str,
    size: float = 10.5,
    bold: bool = False,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    space_after: float = 6,
):
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(space_after)
    run = paragraph.add_run(text)
    set_run_font(run, size, bold)
    return paragraph


def add_heading(
    document,
    text: str,
    size: float = 13,
    color: RGBColor = RGBColor(31, 78, 121),
    space_before: float = 8,
):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_run_font(run, size, True)
    run.font.color.rgb = color
    return paragraph
