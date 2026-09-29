#!/usr/bin/env python3
"""Add submission figures to the W3.1 text for a figure-inclusive review PDF."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph

from build_w3_submission import FIGURE_MAP


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-docx", required=True)
    parser.add_argument("--figure-dir", required=True)
    parser.add_argument("--output-docx", required=True)
    args = parser.parse_args()

    document = Document(Path(args.source_docx))
    figure_dir = Path(args.figure_dir)
    for number, stem in enumerate(FIGURE_MAP.values(), start=1):
        image = figure_dir / f"{stem}.png"
        if not image.is_file():
            raise SystemExit(f"missing review image: {image}")
        legend = next(
            (
                paragraph
                for paragraph in document.paragraphs
                if paragraph.text.strip().startswith(f"Figure {number}.")
            ),
            None,
        )
        if legend is None:
            raise SystemExit(f"missing legend for Figure {number}")
        element = OxmlElement("w:p")
        legend._p.addnext(element)
        paragraph = Paragraph(element, legend._parent)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(image), width=Inches(6.2))

    output = Path(args.output_docx)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)


if __name__ == "__main__":
    main()
