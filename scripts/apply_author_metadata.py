#!/usr/bin/env python3
"""Fill author-only fields and build the BMC cover letter."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


REQUIRED_PATHS = (
    ("affiliations",),
    ("authors",),
    ("authors_contributions",),
    ("competing_interests",),
    ("cover_letter", "date"),
    ("cover_letter", "signatory"),
    ("funding",),
    ("repository_url",),
    ("zenodo_doi",),
)


def get_nested(data: dict, path: tuple[str, ...]):
    value = data
    for key in path:
        value = value[key]
    return value


def missing_required(data: dict) -> list[str]:
    missing: list[str] = []
    for path in REQUIRED_PATHS:
        try:
            value = get_nested(data, path)
        except (KeyError, TypeError):
            missing.append(".".join(path))
            continue
        if value in ("", None, [], {}):
            missing.append(".".join(path))
    authors = data.get("authors", [])
    if not authors or not any(
        author.get("full_name", "").strip() for author in authors
    ):
        missing.append("authors[].full_name")
    if not any(
        author.get("corresponding", False)
        and author.get("email", "").strip()
        for author in authors
    ):
        missing.append("corresponding author email")
    if not data.get("confirmations", {}).get("all_authors_approved"):
        missing.append("confirmations.all_authors_approved")
    if not data.get("confirmations", {}).get("not_published_elsewhere"):
        missing.append("confirmations.not_published_elsewhere")
    return sorted(set(missing))


def set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = text


def find_paragraph(document: Document, prefix: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise RuntimeError(f"paragraph not found: {prefix}")


def format_authors(data: dict) -> str:
    parts = []
    for author in data["authors"]:
        affiliations = ", ".join(author.get("affiliations", []))
        suffix = f" ({affiliations})" if affiliations else ""
        parts.append(author["full_name"].strip() + suffix)
    return ", ".join(parts)


def format_affiliations(data: dict) -> str:
    parts = []
    for affiliation in data["affiliations"]:
        bits = [
            affiliation.get("name", "").strip(),
            affiliation.get("department", "").strip(),
            affiliation.get("city", "").strip(),
            affiliation.get("postal_code", "").strip(),
            affiliation.get("country", "").strip(),
        ]
        bits = [bit for bit in bits if bit]
        parts.append(f"{affiliation['id']}. " + ", ".join(bits))
    return "; ".join(parts)


def format_funding(data: dict) -> str:
    if not data.get("funding"):
        return "The authors received no specific funding for this work."
    parts = []
    for item in data["funding"]:
        funder = item.get("funder", "").strip()
        grant = item.get("grant_id", "").strip()
        parts.append(f"{funder} ({grant})" if grant else funder)
    return "This work was supported by " + "; ".join(parts) + "."


def fill_manuscript(document: Document, data: dict) -> None:
    title = data["title"].strip()
    for paragraph in document.paragraphs:
        if paragraph.style and paragraph.style.name == "Title":
            set_paragraph_text(paragraph, title)
            break

    set_paragraph_text(find_paragraph(document, "[Author names]"), format_authors(data))
    set_paragraph_text(
        find_paragraph(document, "[Affiliations]"),
        "Affiliations: " + format_affiliations(data),
    )
    corresponding = next(
        author for author in data["authors"] if author.get("corresponding")
    )
    set_paragraph_text(
        find_paragraph(document, "Corresponding author:"),
        f"Corresponding author: {corresponding['full_name']}, "
        f"{corresponding.get('email', '').strip()}",
    )

    availability = (
        "Availability of data and materials: The datasets supporting the "
        f"conclusions of this article are available in the repository at "
        f"{data['repository_url']} and archived at "
        f"https://doi.org/{data['zenodo_doi']}."
    )
    for needle, replacement in (
        ("Availability of data and materials:", availability),
        (
            "Competing interests:",
            "Competing interests: " + data["competing_interests"].strip(),
        ),
        ("Funding:", format_funding(data)),
        (
            "Authors' contributions:",
            "Authors' contributions: "
            + data["authors_contributions"].strip(),
        ),
        (
            "Acknowledgements:",
            "Acknowledgements: "
            + (data.get("acknowledgements", "").strip() or "Not applicable."),
        ),
    ):
        set_paragraph_text(find_paragraph(document, needle), replacement)


def add_body(document: Document, text: str, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold


def build_cover_letter(data: dict, output: Path) -> None:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    title.add_run(
        data["cover_letter"].get("date", "").strip() or "[DATE]"
    ).bold = False
    document.add_paragraph(data["cover_letter"]["editor"])
    document.add_paragraph(data["cover_letter"]["journal"])
    document.add_paragraph(
        "Re: Submission of a Research article to the collection "
        f"\"{data['cover_letter']['collection']}\""
    )
    add_body(
        document,
        f"Dear {data['cover_letter']['editor']},",
    )
    add_body(
        document,
        "We are pleased to submit our manuscript entitled "
        f"\"{data['title']}\" for consideration as a Research article in "
        "BMC Bioinformatics.",
    )
    add_body(
        document,
        "The manuscript presents GCerrHMM, a GC-aware error hidden Markov "
        "model trained from real long-read alignments, and a reproducible "
        "fidelity harness. The paper reports a pre-registered, contracted "
        "GC-conditional claim, a six-genome directional panel, and "
        "coverage-matched downstream consistency that ranks the model "
        "second on alignment identity and first on assembly Borda distance "
        "among the compared tools.",
    )
    add_body(
        document,
        "We believe the work fits the collection because it treats simulator "
        "fidelity as a measurement problem: the manuscript separates "
        "descriptive composite scores from the GC-stratified instrument, "
        "reports the noise floor, and releases the fail-loud evaluation "
        "harness with the model.",
    )
    add_body(
        document,
        "The manuscript has not been published or submitted for publication "
        "elsewhere, and all authors have approved this submission."
        if data.get("confirmations", {}).get("not_published_elsewhere")
        and data.get("confirmations", {}).get("all_authors_approved")
        else "AUTHOR CONFIRMATION REQUIRED: publication status and author "
        "approval must be confirmed before submission.",
    )
    add_body(
        document,
        "Competing interests: "
        + (
            data["competing_interests"].strip()
            or "[COMPETING INTERESTS TO BE SUPPLIED]"
        ),
    )
    repository_url = data.get("repository_url", "").strip() or "[REPOSITORY URL]"
    zenodo_doi = data.get("zenodo_doi", "").strip() or "[ZENODO DOI]"
    add_body(
        document,
        "We confirm that the data and code used in this study are available "
        f"at {repository_url} and that the archived release is "
        f"available at https://doi.org/{zenodo_doi}.",
    )
    add_body(document, "Yours sincerely,")
    add_body(
        document,
        data["cover_letter"].get("signatory", "").strip()
        or "[CORRESPONDING AUTHOR]",
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--w3-docx", required=True)
    parser.add_argument("--metadata", required=True)
    parser.add_argument("--output-docx")
    parser.add_argument("--output-cover-letter", required=True)
    parser.add_argument("--allow-incomplete", action="store_true")
    args = parser.parse_args()

    metadata_path = Path(args.metadata)
    data = json.loads(metadata_path.read_text(encoding="utf-8"))
    missing = missing_required(data)
    if missing and not args.allow_incomplete:
        raise SystemExit(
            "author metadata is incomplete: " + ", ".join(missing)
        )

    if args.output_docx:
        document = Document(Path(args.w3_docx))
        fill_manuscript(document, data)
        output = Path(args.output_docx)
        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(output)
        print(f"AUTHOR_METADATA_DOCX={output}")

    build_cover_letter(data, Path(args.output_cover_letter))
    print(f"COVER_LETTER={args.output_cover_letter}")


if __name__ == "__main__":
    main()
